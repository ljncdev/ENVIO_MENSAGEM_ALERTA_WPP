import argparse
import csv
import io
import json
import re
import sys
import time
import unicodedata
import webbrowser
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time as datetime_time, timedelta
from itertools import chain
from pathlib import Path
from typing import Iterable, Optional


SCHEDULE_TYPES = (
    "ENTRADA",
    "SAÍDA ALMOÇO",
    "VOLTA ALMOÇO",
    "SAÍDA",
)
REMINDER_MESSAGES = {
    "ENTRADA": (
        "Olá, {nome}! Já são {horario}. "
        "Não se esqueça de registrar seu ponto de entrada."
    ),
    "SAÍDA ALMOÇO": (
        "Olá, {nome}! Já são {horario}. "
        "Não se esqueça de registrar sua saída para o almoço."
    ),
    "VOLTA ALMOÇO": (
        "Olá, {nome}! Já são {horario}. "
        "O intervalo terminou; registre seu retorno do almoço."
    ),
    "SAÍDA": (
        "Olá, {nome}! Já são {horario}. "
        "Seu expediente está terminando; não se esqueça de registrar sua saída."
    ),
}
POLL_INTERVAL_SECONDS = 1
WHATSAPP_WAIT_SECONDS = 15


@dataclass(frozen=True)
class Employee:
    name: str
    phone: str
    schedule: tuple[Optional[datetime_time], ...]


def normalize_header(value: str) -> str:
    without_accents = unicodedata.normalize("NFKD", value.strip())
    return "".join(
        character
        for character in without_accents
        if not unicodedata.combining(character)
    ).lower()


def row_is_header(row: list[str]) -> bool:
    if len(row) < 2:
        return False
    name_column = normalize_header(row[0])
    phone_column = normalize_header(row[1])
    name_header = name_column.startswith(("nome", "funcionario"))
    phone_header = phone_column.startswith(("telefone", "celular", "whatsapp"))
    return name_header and phone_header


def parse_phone(value: str, row_number: int) -> str:
    phone = re.sub(r"[\s()-]", "", value.strip())
    if not re.fullmatch(r"\+?[1-9]\d{7,14}", phone):
        raise ValueError(
            f"Linha {row_number}: telefone inválido. "
            "Use o código do país, por exemplo +5511999999999."
        )
    return phone if phone.startswith("+") else f"+{phone}"


def parse_schedule_time(value: str, row_number: int, event_type: str) -> Optional[datetime_time]:
    value = value.strip()
    if not value:
        return None
    if not re.fullmatch(r"\d{2}:\d{2}", value):
        raise ValueError(
            f"Linha {row_number}: horário inválido para {event_type}: {value!r}. "
            "Use o formato HH:MM."
        )
    try:
        return datetime.strptime(value, "%H:%M").time()
    except ValueError as error:
        raise ValueError(
            f"Linha {row_number}: horário inválido para {event_type}: {value!r}."
        ) from error


def excel_cell_to_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime_time):
        return value.strftime("%H:%M")
    return str(value)


def load_excel_rows(excel_path: Path) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "Dependência ausente para ler planilhas Excel. "
            "Instale com: pip install -r requirements.txt"
        ) from error

    with excel_path.open("rb") as excel_file:
        workbook = load_workbook(excel_file, read_only=True, data_only=True)
        try:
            worksheet = workbook.active
            rows = [
                [excel_cell_to_text(value) for value in row]
                for row in worksheet.iter_rows(values_only=True)
            ]
            if worksheet.max_column == 1:
                content = "\n".join(row[0] for row in rows if row and row[0].strip())
                if ";" in content or "," in content:
                    sample = content[:4096]
                    try:
                        delimiter = csv.Sniffer().sniff(
                            sample, delimiters=";,"
                        ).delimiter
                    except csv.Error:
                        delimiter = (
                            ";" if sample.count(";") > sample.count(",") else ","
                        )
                    return list(
                        csv.reader(io.StringIO(content), delimiter=delimiter)
                    )
            return rows
        finally:
            workbook.close()


def load_employees_from_rows(
    rows: Iterable[list[str]], csv_path: Path
) -> list[Employee]:
    row_iterator = iter(rows)
    first_row = next(row_iterator, None)
    if first_row is None:
        raise ValueError(f"O arquivo CSV está vazio: {csv_path}")

    has_header = row_is_header(first_row)
    employee_rows = (
        enumerate(row_iterator, start=2)
        if has_header
        else enumerate(chain((first_row,), row_iterator), start=1)
    )
    employees: list[Employee] = []

    for row_number, row in employee_rows:
        if not row or all(not cell.strip() for cell in row):
            continue
        if len(row) != 6:
            raise ValueError(
                f"Linha {row_number}: esperado 6 colunas "
                "(nome, telefone e 4 horários), mas foram encontradas "
                f"{len(row)}."
            )

        name = row[0].strip()
        if not name:
            raise ValueError(f"Linha {row_number}: nome do funcionário vazio.")
        phone = parse_phone(row[1], row_number)
        schedule = tuple(
            parse_schedule_time(row[index + 2], row_number, event_type)
            for index, event_type in enumerate(SCHEDULE_TYPES)
        )
        employees.append(
            Employee(
                name=name,
                phone=phone,
                schedule=schedule,
            )
        )

    if not employees:
        raise ValueError(f"Nenhum funcionário foi encontrado no CSV: {csv_path}")
    return employees


def load_employees(csv_path: Path) -> list[Employee]:
    if not csv_path.is_file():
        raise FileNotFoundError(f"Arquivo CSV não encontrado: {csv_path}")

    if zipfile.is_zipfile(csv_path):
        return load_employees_from_rows(load_excel_rows(csv_path), csv_path)

    with csv_path.open("r", encoding="utf-8-sig", newline="") as csv_file:
        sample = csv_file.read(4096)
        if not sample.strip():
            raise ValueError(f"O arquivo CSV está vazio: {csv_path}")
        csv_file.seek(0)
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=";,").delimiter
        except csv.Error:
            delimiter = ";" if sample.count(";") > sample.count(",") else ","

        return load_employees_from_rows(
            csv.reader(csv_file, delimiter=delimiter),
            csv_path,
        )


def load_sent_events(state_path: Path) -> dict[str, set[str]]:
    if not state_path.exists():
        return {}

    with state_path.open("r", encoding="utf-8") as state_file:
        saved_state = json.load(state_file)
    if not isinstance(saved_state, dict):
        raise ValueError(f"Estado inválido em {state_path}: esperado um objeto JSON.")

    sent_events: dict[str, set[str]] = {}
    for day, events in saved_state.items():
        try:
            date.fromisoformat(day)
        except (TypeError, ValueError) as error:
            raise ValueError(f"Estado inválido em {state_path}: data inválida.") from error
        if not isinstance(events, list) or not all(
            isinstance(event, str) for event in events
        ):
            raise ValueError(f"Estado inválido em {state_path}: lista de envios inválida.")
        sent_events[day] = set(events)
    return sent_events


def save_sent_events(state_path: Path, sent_events: dict[str, set[str]], today: date) -> None:
    cutoff = (today - timedelta(days=7)).isoformat()
    state_path.parent.mkdir(parents=True, exist_ok=True)
    state_to_save = {
        day: sorted(events)
        for day, events in sent_events.items()
        if day >= cutoff
    }
    temporary_path = state_path.with_name(f"{state_path.name}.tmp")
    with temporary_path.open("w", encoding="utf-8") as state_file:
        json.dump(state_to_save, state_file, ensure_ascii=False, indent=2)
        state_file.write("\n")
    temporary_path.replace(state_path)


def event_key(employee: Employee, event_type: str, scheduled_time: datetime_time) -> str:
    return (
        f"{employee.phone}|{employee.name.casefold()}|{event_type}|"
        f"{scheduled_time.strftime('%H:%M')}"
    )


def send_whatsapp_message(phone: str, message: str) -> None:
    from urllib.parse import quote

    try:
        import pyautogui
    except ImportError as error:
        raise RuntimeError(
            "Dependência ausente. Instale com: pip install -r requirements.txt"
        ) from error

    url = (
        f"https://web.whatsapp.com/send?phone={quote(phone, safe='')}"
        f"&text={quote(message)}"
    )
    if not webbrowser.open(url):
        raise RuntimeError("Não foi possível abrir o WhatsApp Web no navegador padrão.")

    if sys.platform == "win32":
        pyautogui.hotkey("win", "up")
    time.sleep(WHATSAPP_WAIT_SECONDS)

    screen_width, screen_height = pyautogui.size()
    pyautogui.click(screen_width - 50, screen_height - 80)
    time.sleep(3)
    pyautogui.hotkey("command" if sys.platform == "darwin" else "ctrl", "w")


def run_scheduler(employees: list[Employee], state_path: Path) -> None:
    sent_events = load_sent_events(state_path)
    print(
        f"Automação iniciada para {len(employees)} funcionário(s). "
        "Pressione Ctrl+C para encerrar."
    )

    while True:
        now = datetime.now().astimezone()
        today_key = now.date().isoformat()
        current_time = now.time().replace(second=0, microsecond=0)
        due_events = []

        for employee in employees:
            for event_type, scheduled_time in zip(SCHEDULE_TYPES, employee.schedule):
                if scheduled_time is None or scheduled_time != current_time:
                    continue
                key = event_key(employee, event_type, scheduled_time)
                if key not in sent_events.get(today_key, set()):
                    due_events.append((employee, event_type, scheduled_time, key))

        for employee, event_type, scheduled_time, key in due_events:
            message = REMINDER_MESSAGES[event_type].format(
                nome=employee.name,
                horario=scheduled_time.strftime("%H:%M"),
            )
            print(
                f"Enviando lembrete de {event_type.lower()} "
                f"para {employee.name} ({scheduled_time.strftime('%H:%M')})."
            )
            send_whatsapp_message(employee.phone, message)
            sent_events.setdefault(today_key, set()).add(key)
            save_sent_events(state_path, sent_events, now.date())

        time.sleep(POLL_INTERVAL_SECONDS)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Envia lembretes de ponto pelo WhatsApp Web conforme um CSV."
    )
    parser.add_argument(
        "--csv",
        type=Path,
        default=Path("funcionarios.csv"),
        help="Caminho do CSV (padrão: funcionarios.csv).",
    )
    parser.add_argument(
        "--state-file",
        type=Path,
        default=Path("estado_envios.json"),
        help="Arquivo que registra lembretes já enviados.",
    )
    parser.add_argument(
        "--check",
        action="store_true",
        help="Valida o CSV e encerra sem enviar mensagens.",
    )
    args = parser.parse_args()

    employees = load_employees(args.csv)
    if args.check:
        print(f"CSV válido: {len(employees)} funcionário(s) carregado(s).")
        return

    run_scheduler(employees, args.state_file)


if __name__ == "__main__":
    main()
